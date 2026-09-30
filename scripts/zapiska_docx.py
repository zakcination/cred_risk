# -*- coding: utf-8 -*-
"""Сборка документа для руководства в формате записки v4 — `docs/STYLE_ZAPISKA.md`.

Источник — текстовый файл с простой разметкой, результат — .docx:

    ---
    title: ОТЧЁТ
    subtitle: по результатам ...
    year: 2026 год
    ---
    ## 1. Цель и основания            заголовок первого уровня
    ### 5.1. Изменение базы расчёта   заголовок второго уровня
    Абзац текста. **полужирный** фрагмент.
    **Весь абзац полужирным — ключевая фраза.**
    - пункт перечисления               маркер «•»
    1. пункт вывода или предложения    нумерация как в тексте
    = M(T) = z(p) × σΔ × √T            формула: по центру, полужирным
    Таблица 1. Подпись                 подпись над таблицей, дальше — таблица
    | Шапка | ... |
    |---|---|
    | ячейка | ... |
    Примечание: ...                    мелким кеглем
    \\pagebreak                        разрыв страницы

Оформление — по разделу 1 руководства: A4, поля 2,0 / 2,0 / 1,75 / 1,5 см,
Times New Roman 12, по ширине, отступ первой строки 1,25 см, колонтитул
«ВНУТРЕННЯЯ ИНФОРМАЦИЯ», титульный лист.

Собранный .docx в репозиторий не коммитится: колонтитул помечает его как
внутреннюю информацию (§2 корневого CLAUDE.md). Коммитится источник.

    python3 scripts/zapiska_docx.py <источник.md> <результат.docx>
    python3 scripts/zapiska_docx.py --selftest
"""

import os
import re
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"
MARK = "ВНУТРЕННЯЯ ИНФОРМАЦИЯ"
BODY_PT = 12
NOTE_PT = 10
INDENT = Cm(1.25)
LIST_LEFT = Cm(1.27)
LIST_HANG = Cm(0.63)

NUM_CELL = re.compile(r"^[\s\d\-−–+,.%()≈<>≤≥×/п₸млнрдтыс]*$")


def set_font(run, size=BODY_PT, bold=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if bold is not None:
        run.bold = bold


def add_runs(par, text, size=BODY_PT, bold_all=False):
    """Текст с разметкой **полужирного**."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(par.add_run(part[2:-2]), size, True)
        else:
            set_font(par.add_run(part), size, True if bold_all else None)


def fmt(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=INDENT, before=0, after=6,
        left=None, hanging=None, keep=False):
    pf = par.paragraph_format
    par.alignment = align
    pf.first_line_indent = first
    if left is not None:
        pf.left_indent = left
    if hanging is not None:
        pf.first_line_indent = -hanging
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.15
    pf.keep_with_next = keep


def base_document():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(1.75), Cm(1.5)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(hp.add_run(MARK), 10, name="Calibri")
    return doc


def title_page(doc, meta):
    for _ in range(8):
        fmt(doc.add_paragraph(), first=None, after=0)
    p = doc.add_paragraph()
    fmt(p, WD_ALIGN_PARAGRAPH.CENTER, None, after=6)
    set_font(p.add_run(meta.get("title", "")), 14, True)
    if meta.get("subtitle"):
        p = doc.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, None, after=0)
        set_font(p.add_run(meta["subtitle"]), BODY_PT, True)
    for _ in range(20):
        fmt(doc.add_paragraph(), first=None, after=0)
    p = doc.add_paragraph()
    fmt(p, WD_ALIGN_PARAGRAPH.CENTER, None, after=0)
    set_font(p.add_run(meta.get("year", "")), BODY_PT, True)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    size = 12 if ncol <= 3 else 11 if ncol <= 5 else 10
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncol):
            text = row[j] if j < len(row) else ""
            cell = t.cell(i, j)
            par = cell.paragraphs[0]
            numeric = i > 0 and j > 0 and text and NUM_CELL.match(text)
            fmt(par, WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or numeric)
                else WD_ALIGN_PARAGRAPH.LEFT, None, after=0)
            add_runs(par, text, size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse(src):
    """(meta, blocks). Блок — (вид, содержимое)."""
    lines = src.splitlines()
    meta, i = {}, 0
    if lines and lines[0].strip() == "---":
        i = 1
        while lines[i].strip() != "---":
            k, _, v = lines[i].partition(":")
            meta[k.strip()] = v.strip()
            i += 1
        i += 1
    blocks, para = [], []

    def flush():
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    while i < len(lines):
        s = lines[i].rstrip()
        st = s.strip()
        if not st:
            flush()
        elif st == "\\pagebreak":
            flush(); blocks.append(("break", ""))
        elif st.startswith("### "):
            flush(); blocks.append(("h2", st[4:]))
        elif st.startswith("## "):
            flush(); blocks.append(("h1", st[3:]))
        elif st.startswith("= "):
            flush(); blocks.append(("formula", st[2:]))
        elif st.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = split_row(lines[i])
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in row):
                    rows.append(row)
                i += 1
            blocks.append(("table", rows))
            continue
        elif re.match(r"^Таблица \d+\.", st):
            flush(); blocks.append(("caption", st))
        elif st.startswith("- "):
            flush(); blocks.append(("bullet", st[2:]))
        elif re.match(r"^\d+\.\s", st) and not para:
            num, _, rest = st.partition(" ")
            blocks.append(("num", (num, rest)))
        elif (st.startswith("  ") or s.startswith("  ")) and blocks and \
                blocks[-1][0] in ("bullet", "num") and not para:
            kind, val = blocks[-1]
            if kind == "bullet":
                blocks[-1] = ("bullet", val + " " + st)
            else:
                blocks[-1] = ("num", (val[0], val[1] + " " + st))
        else:
            para.append(st)
        i += 1
    flush()
    return meta, blocks


def build(src_path, out_path):
    with open(src_path, encoding="utf-8") as fh:
        meta, blocks = parse(fh.read())
    doc = base_document()
    title_page(doc, meta)
    n_tables = 0
    for kind, val in blocks:
        if kind == "h1":
            p = doc.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.LEFT, None, before=12, after=6, keep=True)
            set_font(p.add_run(val), BODY_PT, True)
        elif kind == "h2":
            p = doc.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.LEFT, None, before=6, after=6, keep=True)
            set_font(p.add_run(val), BODY_PT, True)
        elif kind == "caption":
            p = doc.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.LEFT, None, before=6, after=4, keep=True)
            set_font(p.add_run(val), BODY_PT, True)
        elif kind == "table":
            add_table(doc, val)
            n_tables += 1
        elif kind == "formula":
            p = doc.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.CENTER, None, before=4, after=8)
            set_font(p.add_run(val), BODY_PT, True)
        elif kind == "bullet":
            p = doc.add_paragraph()
            fmt(p, left=LIST_LEFT, hanging=LIST_HANG, after=3)
            set_font(p.add_run("•\t"), BODY_PT)
            add_runs(p, val)
        elif kind == "num":
            p = doc.add_paragraph()
            fmt(p, left=LIST_LEFT, hanging=LIST_HANG, after=4)
            set_font(p.add_run(val[0] + "\t"), BODY_PT)
            add_runs(p, val[1])
        elif kind == "break":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            p = doc.add_paragraph()
            if val.startswith("Примечание"):
                fmt(p, first=None, after=6)
                add_runs(p, val, NOTE_PT)
            elif val.startswith("**") and val.endswith("**") and val.count("**") == 2:
                fmt(p)
                add_runs(p, val[2:-2], bold_all=True)
            else:
                fmt(p)
                add_runs(p, val)
    doc.save(out_path)
    return len(blocks), n_tables


def selftest():
    ok = True

    def check(name, cond, detail=""):
        nonlocal ok
        ok &= bool(cond)
        print(f"  [{'ok' if cond else 'СБОЙ'}] {name}{'  ' + detail if detail else ''}")

    sample = """---
title: ОТЧЁТ
subtitle: по результатам проверки
year: 2026 год
---
## 1. Цель
Первая строка абзаца
продолжение абзаца с **выделением**.

**Таким образом, это ключевая фраза.**

- пункт один;
- пункт два
  с переносом.

= M(T) = z × σ × √T

Таблица 1. Подпись
| Показатель | Значение |
|---|---:|
| Первый | 12,5% |

Примечание: мелко.

### 1.1. Подраздел
1. Вывод первый.
2. Вывод второй.
"""
    meta, blocks = parse(sample)
    kinds = [k for k, _ in blocks]
    check("шапка разобрана", meta == {"title": "ОТЧЁТ", "subtitle": "по результатам проверки",
                                       "year": "2026 год"}, str(meta))
    check("строки абзаца склеены", ("p", "Первая строка абзаца продолжение абзаца с **выделением**.")
          in blocks)
    check("перенос в пункте списка склеен", ("bullet", "пункт два с переносом.") in blocks)
    check("подпись стоит перед таблицей", kinds.index("caption") + 1 == kinds.index("table"))
    check("строка-разделитель таблицы отброшена",
          [b for b in blocks if b[0] == "table"][0][1] == [["Показатель", "Значение"],
                                                           ["Первый", "12,5%"]])
    check("нумерованные пункты распознаны", kinds.count("num") == 2)
    with tempfile.TemporaryDirectory() as tmp:
        src, out = os.path.join(tmp, "a.md"), os.path.join(tmp, "a.docx")
        with open(src, "w", encoding="utf-8") as fh:
            fh.write(sample)
        build(src, out)
        d = Document(out)
        sec = d.sections[0]
        check("колонтитул «ВНУТРЕННЯЯ ИНФОРМАЦИЯ»", sec.header.paragraphs[0].text == MARK)
        check("поля 2,0 / 2,0 / 1,75 / 1,5 см",
              [round(x.cm, 2) for x in (sec.top_margin, sec.bottom_margin,
                                         sec.left_margin, sec.right_margin)]
              == [2.0, 2.0, 1.75, 1.5])
        texts = [p.text for p in d.paragraphs]
        check("титул первым", "ОТЧЁТ" in texts[:12])
        body = [p for p in d.paragraphs if p.text.startswith("Первая строка")][0]
        check("абзац по ширине с отступом 1,25 см",
              body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
              and round(body.paragraph_format.first_line_indent.cm, 2) == 1.25)
        key = [p for p in d.paragraphs if p.text.startswith("Таким образом")][0]
        check("ключевая фраза полужирная целиком", all(r.bold for r in key.runs))
        f = [p for p in d.paragraphs if p.text.startswith("M(T)")][0]
        check("формула по центру", f.alignment == WD_ALIGN_PARAGRAPH.CENTER)
        check("шрифт Times New Roman в тексте",
              all(r.font.name == FONT for r in body.runs))
        check("одна таблица, стиль сетки", len(d.tables) == 1
              and d.tables[0].style.name == "Table Grid")
    return ok


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        print("Самопроверка zapiska_docx:")
        sys.exit(0 if selftest() else 1)
    if len(sys.argv) != 3:
        sys.exit("нужно: zapiska_docx.py <источник.md> <результат.docx>")
    n, t = build(sys.argv[1], sys.argv[2])
    print(f"{sys.argv[2]}: собрано, блоков {n}, таблиц {t}")
