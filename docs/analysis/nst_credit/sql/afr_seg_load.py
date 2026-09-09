# -*- coding: utf-8 -*-
"""Распаковка эталонной сегментации АФР из защищённой книги в CSV для загрузки в БД.

Зачем отдельный шаг. Файл АФР приходит зашифрованным (OLE-контейнер с потоками
EncryptionInfo / EncryptedPackage), openpyxl его не открывает вовсе — падает
с «File is not a zip file». Расшифровка выполняется здесь, пароль передаётся
аргументом и в код не зашивается.

Что делает:
  1. расшифровывает книгу во временный файл;
  2. выгружает лист в CSV с разделителем «;» и кодировкой UTF-8 без BOM —
     готовый вход для BULK INSERT;
  3. печатает профиль: значения сегмента, уникальность ключа, расхождения
     между тремя колонками идентификатора.

Профиль печатается всегда: без него неизвестно, тот ли это файл, и совпадает ли
он с таблицей эталона в базе. Числа профиля сверяются выводом 0 скрипта
seg_2024_vs_afr.sql.

§ 2: выгруженный CSV остаётся на машине автора, в репозиторий не коммитится.

Запуск:
  python afr_seg_load.py EUB_SEGMENTS.xlsx --password "..." --out afr_segments.csv
"""

import argparse
import csv
import os
import sys
import tempfile
from collections import Counter


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path", help="защищённая книга АФР")
    ap.add_argument("--password", required=True)
    ap.add_argument("--out", default="afr_segments.csv")
    ap.add_argument("--sheet", default=None, help="имя листа; по умолчанию первый")
    a = ap.parse_args()

    if not os.path.isfile(a.path):
        sys.exit("файл не найден: %s" % a.path)

    try:
        import msoffcrypto
    except ImportError:
        sys.exit("нет библиотеки msoffcrypto-tool: pip install msoffcrypto-tool")
    from openpyxl import load_workbook

    tmpdir = tempfile.mkdtemp(prefix="afrseg_")
    dec = os.path.join(tmpdir, "decrypted.xlsx")
    with open(a.path, "rb") as f:
        off = msoffcrypto.OfficeFile(f)
        off.load_key(password=a.password)
        with open(dec, "wb") as g:
            off.decrypt(g)
    print("расшифровано во временный файл (%d байт)" % os.path.getsize(dec))

    wb = load_workbook(dec, read_only=True, data_only=True)
    ws = wb[a.sheet] if a.sheet else wb.worksheets[0]
    print("лист «%s»: строк %s, колонок %s" % (ws.title, ws.max_row, ws.max_column))

    rows = ws.iter_rows(values_only=True)
    header = [("" if v is None else str(v).strip()) for v in next(rows)]
    print("заголовок: %s" % " | ".join(header))
    idx = {h.upper(): i for i, h in enumerate(header)}
    for need in ("LOAN_ID_KR", "SEGMENT"):
        if need not in idx:
            sys.exit("в файле нет колонки %s — сверить формат выгрузки" % need)

    seg = Counter()
    keys = Counter()
    diff = Counter()
    n = 0
    with open(a.out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([h.lower() for h in header])
        for r in rows:
            n += 1
            vals = ["" if v is None else str(v).strip() for v in r[:len(header)]]
            w.writerow(vals)
            kr = vals[idx["LOAN_ID_KR"]]
            seg[vals[idx["SEGMENT"]]] += 1
            keys[kr] += 1
            for col in ("LOAN_ID", "ID"):
                if col in idx and vals[idx[col]] != kr:
                    diff[col] += 1
    wb.close()

    print("\nвыгружено строк: %d -> %s" % (n, a.out))
    print("\n=== ЗНАЧЕНИЯ SEGMENT ===")
    for k, v in seg.most_common():
        print("  %-24s %8d  %5.2f %%" % (k or "(пусто)", v, 100.0 * v / n))
    dup = sum(1 for v in keys.values() if v > 1)
    print("\n=== КЛЮЧ LOAN_ID_KR ===")
    print("  уникальных : %d" % len(keys))
    print("  дублей     : %d" % dup)
    print("  пустых     : %d" % keys.get("", 0))
    print("\n=== РАСХОЖДЕНИЯ КОЛОНОК ИДЕНТИФИКАТОРА ===")
    for col in ("LOAN_ID", "ID"):
        if col in idx:
            print("  %-10s != LOAN_ID_KR : %d (%.1f %%)" % (col, diff[col], 100.0 * diff[col] / n))
    print("\nСоединять только по LOAN_ID_KR. Ключ строковый — приводить обе стороны")
    print("к одному типу и не обрезать ведущие нули.")

    try:
        os.remove(dec); os.rmdir(tmpdir)
    except OSError:
        pass


if __name__ == "__main__":
    main()
