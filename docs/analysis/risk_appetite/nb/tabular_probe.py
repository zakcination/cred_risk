# -*- coding: utf-8 -*-
"""
Аудит схемы табличных файлов: что за листы, какие колонки, какого типа, за какой период.

Отличие от limits_probe.py: тот ищет СЛОВА в прозе (ВНД, протоколы) и на числовых
таблицах слеп — дата приходит объектом datetime, доля лежит числом без знака «%».
Этот читает СТРУКТУРУ и ничего не ищет по словам.

§ 4: имена колонок берутся из живого аудита, а не из памяти. Скрипт — этот аудит.
§ 6: схему показывает, ряда не строит, полей не выбирает.
§ 2: собранные CSV остаются на машине автора.

Запуск:
  python tabular_probe.py "C:\\project_mz\\limits"
  python tabular_probe.py "C:\\project_mz\\limits" --dump-nb        # + исходники ячеек .ipynb
  python tabular_probe.py "C:\\project_mz\\limits" --only 03_series --max-cols 40
"""

import argparse
import csv
import io
import json
import os
import re
import sys
from datetime import date, datetime

XL_EXT = {".xlsx", ".xlsm", ".xltx"}
NUM_RE = re.compile(r"^-?\d{1,3}(?:[ \u00a0]\d{3})*(?:[.,]\d+)?$|^-?\d+(?:[.,]\d+)?$")
DATE_FMTS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y",
             "%Y-%m-%d %H:%M:%S", "%d.%m.%y")


def coerce(v):
    """str -> число / дата / строка. Нестроковые значения возвращаются как есть."""
    if v is None:
        return None
    if isinstance(v, (datetime, date, int, float)) and not isinstance(v, bool):
        return v
    s = str(v).strip()
    if not s:
        return None
    if NUM_RE.match(s):
        try:
            return float(s.replace("\u00a0", "").replace(" ", "").replace(",", "."))
        except ValueError:
            pass
    if re.match(r"^\d{1,4}[-./]\d{1,2}[-./]\d{1,4}", s):
        for f in DATE_FMTS:
            try:
                return datetime.strptime(s[:len(f) + 2].strip(), f)
            except ValueError:
                continue
    return s


def fmt(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return ("%.6g" % v)
    return str(v)


def profile(col):
    """Список значений одной колонки -> словарь профиля."""
    nums, dates, strs = [], [], []
    for v in col:
        v = coerce(v)
        if v is None:
            continue
        if isinstance(v, (datetime, date)):
            dates.append(v)
        elif isinstance(v, (int, float)) and not isinstance(v, bool):
            nums.append(float(v))
        else:
            strs.append(str(v))
    n = len(nums) + len(dates) + len(strs)
    kind = "пусто"
    lo = hi = ""
    note = ""
    sample = ""
    if dates and len(dates) >= max(len(nums), len(strs)):
        kind, lo, hi = "дата", fmt(min(dates)), fmt(max(dates))
    elif nums and len(nums) >= len(strs):
        kind, lo, hi = "число", fmt(min(nums)), fmt(max(nums))
        mx = max(abs(x) for x in nums)
        if 0 < mx <= 1.5:
            note = "доля 0..1 — знака % нет"
        elif 1.5 < mx <= 300 and min(nums) >= -300:
            note = "похоже на проценты"
    elif strs:
        kind = "строка"
        uniq = sorted(set(strs), key=len)
        sample = " | ".join(u[:24] for u in uniq[:6])
        if len(uniq) > 6:
            sample += " | …+%d" % (len(uniq) - 6)
        lo = "uniq=%d" % len(uniq)
    return {"n": n, "kind": kind, "min": lo, "max": hi, "note": note, "sample": sample}


def header_row(rows, max_probe=12):
    """Номер строки-заголовка: первая с >=2 непустыми строковыми ячейками."""
    for i, r in enumerate(rows[:max_probe]):
        s = sum(1 for v in r if isinstance(v, str) and v.strip())
        if s >= 2:
            return i
    return 0


def read_xlsx(path, max_rows, max_cols):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(max_row=max_rows, max_col=max_cols,
                                                  values_only=True)]
            yield ws.title, ws.max_row, ws.max_column, rows
    finally:
        wb.close()


def read_csv_file(path, max_rows, max_cols):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            txt = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        return None
    head = txt[:4096]
    delim = ";" if head.count(";") > head.count(",") else ","
    if head.count("\t") > max(head.count(";"), head.count(",")):
        delim = "\t"
    rows = []
    for i, r in enumerate(csv.reader(io.StringIO(txt), delimiter=delim)):
        if i >= max_rows:
            break
        rows.append(r[:max_cols])
    return delim, enc, rows


def dump_json(path, out_lines):
    with open(path, "rb") as f:
        try:
            obj = json.load(f)
        except Exception as e:
            out_lines.append("    не разобран: %s" % e.__class__.__name__)
            return
    if isinstance(obj, dict):
        for k, v in list(obj.items())[:40]:
            if isinstance(v, (dict, list)):
                out_lines.append("    %-32s %s(%d)" % (k[:32], type(v).__name__, len(v)))
            else:
                out_lines.append("    %-32s %s" % (k[:32], str(v)[:60]))
        if len(obj) > 40:
            out_lines.append("    …ещё %d ключей" % (len(obj) - 40))
    elif isinstance(obj, list):
        out_lines.append("    список, %d элементов" % len(obj))
        if obj and isinstance(obj[0], dict):
            out_lines.append("    ключи первого: %s" % ", ".join(list(obj[0])[:20]))


def dump_nb(path, fh):
    with open(path, "rb") as f:
        try:
            nb = json.load(f)
        except Exception:
            return 0
    n = 0
    fh.write("\n" + "=" * 78 + "\n" + path + "\n" + "=" * 78 + "\n")
    for cell in nb.get("cells", []):
        if cell.get("cell_type") != "code":
            continue
        n += 1
        src = "".join(cell.get("source", []))
        fh.write("\n# ---- ячейка %d ----\n%s\n" % (n, src))
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("root")
    ap.add_argument("--out", default=None)
    ap.add_argument("--only", default=None, help="подстрока в имени файла")
    ap.add_argument("--max-rows", type=int, default=5000)
    ap.add_argument("--max-cols", type=int, default=60)
    ap.add_argument("--dump-nb", action="store_true", help="выгрузить исходники ячеек .ipynb")
    a = ap.parse_args()

    root = os.path.abspath(a.root)
    if not os.path.isdir(root):
        sys.exit("нет такой папки: %s" % root)
    out = os.path.abspath(a.out) if a.out else os.path.join(root, "_probe")
    os.makedirs(out, exist_ok=True)

    schema = []
    lines = []
    nb_fh = open(os.path.join(out, "notebooks_src.txt"), "w", encoding="utf-8") if a.dump_nb else None

    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if d != "_probe" and not d.startswith("~")]
        for fn in sorted(filenames):
            if fn.startswith("~$") or fn.startswith("."):
                continue
            if a.only and a.only.lower() not in fn.lower():
                continue
            full = os.path.join(dirpath, fn)
            rel = os.path.relpath(full, root)
            ext = os.path.splitext(fn)[1].lower()

            if ext == ".ipynb" and nb_fh:
                k = dump_nb(full, nb_fh)
                lines.append("\n### %s  — ноутбук, кодовых ячеек: %d (исходники в notebooks_src.txt)" % (rel, k))
                continue
            if ext == ".json":
                lines.append("\n### %s  — json" % rel)
                dump_json(full, lines)
                continue

            sheets = []
            try:
                if ext in XL_EXT:
                    for title, nr, nc, rows in read_xlsx(full, a.max_rows, a.max_cols):
                        sheets.append((title, nr, nc, rows, ""))
                elif ext == ".csv":
                    got = read_csv_file(full, a.max_rows, a.max_cols)
                    if got is None:
                        lines.append("\n### %s — не декодирован" % rel)
                        continue
                    delim, enc, rows = got
                    sheets.append(("(csv)", len(rows), max((len(r) for r in rows), default=0),
                                   rows, "разделитель «%s», кодировка %s" % (delim, enc)))
                else:
                    continue
            except Exception as e:
                lines.append("\n### %s — ошибка чтения: %s" % (rel, e.__class__.__name__))
                continue

            for title, nr, nc, rows, extra in sheets:
                if not rows:
                    lines.append("\n### %s :: %s — пусто" % (rel, title))
                    continue
                h = header_row(rows)
                names = [str(v).strip() if v is not None else "" for v in rows[h]]
                body = rows[h + 1:]
                lines.append("\n### %s :: %s   строк=%s кол=%s  заголовок@r%d %s"
                             % (rel, title, nr, nc, h + 1, extra))
                for j in range(len(names)):
                    col = [r[j] if j < len(r) else None for r in body]
                    p = profile(col)
                    if p["n"] == 0 and not names[j]:
                        continue
                    nm = names[j] or "(без имени)"
                    rng = ("%s..%s" % (p["min"], p["max"])) if p["max"] else p["min"]
                    lines.append("  c%02d %-28s n=%-6d %-7s %-26s %s%s"
                                 % (j + 1, nm[:28], p["n"], p["kind"], rng[:26],
                                    p["note"], (" :: " + p["sample"]) if p["sample"] else ""))
                    schema.append({"file": rel, "sheet": title, "col_index": j + 1,
                                   "col_name": nm, "n": p["n"], "kind": p["kind"],
                                   "min": p["min"], "max": p["max"], "note": p["note"],
                                   "sample": p["sample"]})

    if nb_fh:
        nb_fh.close()

    sch_p = os.path.join(out, "schema.csv")
    with open(sch_p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["file", "sheet", "col_index", "col_name", "n",
                                          "kind", "min", "max", "note", "sample"],
                           delimiter=";")
        w.writeheader()
        w.writerows(schema)

    print("корень : %s" % root)
    print("колонок в схеме: %d" % len(schema))
    print("вывод  : %s" % sch_p)
    if nb_fh:
        print("         %s" % os.path.join(out, "notebooks_src.txt"))
    print("\n".join(lines))


if __name__ == "__main__":
    main()
