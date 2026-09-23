# -*- coding: utf-8 -*-
"""Сборка .docx из markdown контура — для рассылки коллегам.

Зачем отдельный конвертер. Документы контура ведутся в markdown: он диффится,
и правка видна в PR. Но markdown нельзя отправить коллеге, который откроет его
в Word. Поэтому источником истины остаётся .md, а .docx **всегда пересобирается**
из него и никогда не правится руками — иначе через две итерации они разойдутся,
и будет неизвестно, какой из них верный.

Поддерживается ровно то, чем пользуются документы контура: заголовки H1–H3,
абзацы с **жирным** и `кодом`, таблицы, маркированные и нумерованные списки,
цитаты, блоки кода, горизонтальные линии. Ничего сверх этого сознательно:
конвертер, который умеет всё, невозможно проверить.

Запуск:
    python3 md_to_docx.py ../REVISION_RA.md ../out/Reviziya_RA.docx
    python3 md_to_docx.py --selftest
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

INK = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)

# **жирный**, `код`, остальное — обычный текст
TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_runs(par, text):
    """Разбирает inline-разметку в последовательность run-ов."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            r = par.add_run(part)
        r.font.color.rgb = INK
    return par


def clear_par(par):
    """Убирает существующие run-ы. Присваивание par.text = "" оставляет пустой
    run, и первый значащий run оказывается вторым — самопроверка это поймала."""
    for r in list(par.runs):
        r._element.getparent().remove(r._element)
    return par


def split_row(line):
    """Ячейки строки таблицы без крайних разделителей."""
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_sep(line):
    """Строка-разделитель шапки таблицы: |---|:--:|---|"""
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def build(md_path, docx_path):
    with open(md_path, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    i, n_tables, in_code = 0, 0, False
    while i < len(lines):
        ln = lines[i]

        if ln.strip().startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph()
            r = p.add_run(ln)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(0)
            i += 1
            continue

        s = ln.strip()
        if not s:
            i += 1
            continue

        if re.fullmatch(r"-{3,}", s):
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=len(m.group(1)))
            i += 1
            continue

        # таблица: строка с | и следующая — разделитель
        if s.startswith("|") and i + 1 < len(lines) and is_sep(lines[i + 1]):
            head = split_row(s)
            body, j = [], i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                body.append(split_row(lines[j]))
                j += 1
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, txt in zip(t.rows[0].cells, head):
                add_runs(clear_par(c.paragraphs[0]), f"**{txt}**" if txt else "")
            for row in body:
                cells = t.add_row().cells
                for c, txt in zip(cells, row[:len(head)]):
                    add_runs(clear_par(c.paragraphs[0]), txt)
            doc.add_paragraph()
            n_tables += 1
            i = j
            continue

        if s.startswith(">"):
            p = doc.add_paragraph()
            add_runs(p, s.lstrip("> ").strip())
            p.paragraph_format.left_indent = Pt(24)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = MUTED
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.*)", s)
        if m:
            add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue

        m = re.match(r"^\d+\.\s+(.*)", s)
        if m:
            add_runs(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue

        # обычный абзац: склеить до пустой строки
        buf = [s]
        j = i + 1
        while j < len(lines) and lines[j].strip() and not re.match(
                r"^(#{1,3}\s|\||>|[-*]\s|\d+\.\s|```|-{3,}$)", lines[j].strip()):
            buf.append(lines[j].strip())
            j += 1
        p = add_runs(doc.add_paragraph(), " ".join(buf))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i = j

    os.makedirs(os.path.dirname(os.path.abspath(docx_path)), exist_ok=True)
    doc.save(docx_path)
    return n_tables


def selftest():
    """Проверяется разбор, а не внешний вид: что из чего получилось."""
    import tempfile
    src = (
        "# Заголовок\n\n"
        "Абзац с **жирным** и `кодом`, перенесённый\nна две строки.\n\n"
        "## Раздел\n\n"
        "| A | B |\n|---|---|\n| 1 | **2** |\n| 3 | 4 |\n\n"
        "- пункт\n- ещё\n\n1. раз\n2. два\n\n> цитата\n\n---\n\n"
        "```\nкод\n```\n"
    )
    ok = True

    def check(name, cond, detail=""):
        nonlocal ok
        ok &= bool(cond)
        print(f"  [{'ok' if cond else 'СБОЙ'}] {name}{'  ' + detail if detail else ''}")

    with tempfile.TemporaryDirectory() as d:
        md, dx = os.path.join(d, "a.md"), os.path.join(d, "a.docx")
        open(md, "w", encoding="utf-8").write(src)
        n = build(md, dx)
        doc = Document(dx)
        check("таблица распознана одна", n == 1, f"найдено {n}")
        t = doc.tables[0]
        check("таблица 3×2 (шапка + две строки)",
              len(t.rows) == 3 and len(t.columns) == 2,
              f"{len(t.rows)}×{len(t.columns)}")
        check("шапка жирная", t.rows[0].cells[0].paragraphs[0].runs[0].bold)
        check("жирное в ячейке сохранено",
              any(r.bold for r in t.rows[1].cells[1].paragraphs[0].runs))
        txt = "\n".join(p.text for p in doc.paragraphs)
        check("абзац склеен из двух строк", "перенесённый на две строки" in txt)
        check("маркировка ** убрана из текста", "**" not in txt)
        check("обратные кавычки убраны", "`" not in txt)
        check("заголовок есть", any(
            p.style.name.startswith("Heading") and p.text == "Заголовок"
            for p in doc.paragraphs))
        check("горизонтальная линия не стала абзацем", "---" not in txt)
        check("код сохранён", "код" in txt)
    return ok


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        print("Самопроверка md_to_docx:")
        sys.exit(0 if selftest() else 1)
    if len(sys.argv) < 3:
        sys.exit("нужно: md_to_docx.py <источник.md> <результат.docx>")
    k = build(sys.argv[1], sys.argv[2])
    print(f"{sys.argv[2]}: собрано, таблиц {k}")
